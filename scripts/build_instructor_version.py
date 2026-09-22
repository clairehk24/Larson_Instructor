"""Build instructor HTML pages and downloads from the available L1715 manuscripts."""

from __future__ import annotations

from copy import deepcopy
from html import escape, unescape
from pathlib import Path, PureWindowsPath
import re
from shutil import copy2
import sys
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from lxml import etree

from generate_simulation_downloads import (
    accept_tracked_changes_xml,
    finalize_document,
    stamp_footer,
    write_docx,
)


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPTS = ROOT / "assets" / "Manuscripts"
PAGES = ROOT / "pages"
DOWNLOADS = ROOT / "assets" / "downloads"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PRODUCTION_PREFIX = re.compile(r"^<(cn|ct|a|b|c|d|lh|tt|title|txni|tx|fc)>")
HEADING_TAGS = {"a": "h2", "b": "h3", "c": "h4", "d": "h5", "lh": "h3", "tt": "h3", "title": "h2"}
LIST_SECTIONS = {
    "Learning Objectives": ("ol", "learning-objectives-list"),
    "CAATE 2020 Standards": ("ul", "standards-list"),
    "References": ("ol", "references-list"),
}
INLINE_NUMBERED_ITEMS = {
    17: frozenset(
        {
            "Set the stage",
            "Focus the discussion",
            "Establish a common base",
            "Facilitate with skill",
            "Maintain civility and charity",
            "Encourage reflection and make the connection",
        }
    ),
    31: frozenset(
        {
            "Invite individuals familiar with athletics: faculty, graduate students, athletic trainers (ATs), clinical preceptors, and so on. These individuals may offer meaningful insights and relevant questions;",
            "have students adopt roles and ask questions of their peers (they create questions as part of presimulation Activity 4);",
            "use Embedded Personnel (EPs) to portray specific roles; or",
            "any combination of these approaches.",
        }
    ),
}
SIMULATION_LINKS = {
    17: (
        "https://teaching.uoregon.edu/sites/default/files/2021-04/strategies-for-engaging-with-difficult-topics-strong-emotions-and-challenging-moments-in-the-classroom.pdf",
    ),
}
PAGE_TEXT_REPLACEMENTS = {
    29: {
        "other AT.’": "other AT.",
    },
}
SECTION_CONTENT_ADDITIONS = {
    17: {
        "HCP Prebrief": (
            "Complete the above activities and arrive at class ready to discuss and demonstrate. There is no preparation other than the activities. This simulation is called a table-top simulation. This means that each step will be discussed, practiced, and implemented gradually and as a group. In this learning community, we value differences and expect you to foster a safe space to share your unique perspectives in a respectful manner to enrich the classroom discussions. The instructor will also create a space that is safe for unique, respectful, enriching discussions, and viewpoints.",
            "Because of the unique nature of this simulation, we will not use the safety phrase “Time-out, time-out.” If during this simulation experience, you need to leave the simulation for a break, you may leave the room at any time. Please give the instructor a “thumbs up” as you leave to be sure the instructor knows you are okay. If you do not provide “thumbs up”, someone will follow you out of the classroom to verify you do not need assistance. If you do not return",
        ),
    },
}
FORCED_TOP_LEVEL_HEADINGS = {
    22: frozenset({"Standardized Patient Information"}),
}
SOURCE_LIST_NUMBER_RE = re.compile(r"^(\d+\.\s*)")
SIM_RE = re.compile(r"L1715_Sim(\d{2})")
BEGIN_RE = re.compile(r"BEGIN downloadable content")
BUTTON_RE = re.compile(r"Button name:\s*(.*?)(?=<title>|\\?$)")
NOTE_END_RE = re.compile(r"xqq\\$", re.IGNORECASE)
PAGE_BREAK_MARKER_RE = re.compile(r"^\\qqStart new page", re.IGNORECASE)
TITLE_OVERRIDES = {
    5: "Special Test Roulette: Upper Extremity",
    9: "Coach Education",
}
SUPPORTING_RESOURCES = {
    2: (
        "L1715_Sim02_TS 02.01_Rowing-Basics-3.pdf",
        "L1715_Sim02_TS 02.02_Understanding Rowing.htm",
    ),
    16: (
        "L1715_Sim16_Instructor Slides.pptx",
        "L1715_Sim16_Notecards for HCP Evaluation 1.docx",
        "L1715_Sim16_Notecards for HCP Evaluation 2.docx",
    ),
    17: ("L1715_Sim17_Instructor Slides.pptx",),
    21: (
        "L1715_Sim21 HCP SCAT6 Rain Shoemaker.pdf",
        "L1715_Sim21 SP SCAT6 Rain Shoemaker.pdf",
    ),
    22: ("L1715_Sim22 SCAT6 Rain Shoemaker 96 Hours.pdf",),
}
SUPPORTING_OUTPUT_NAMES = {
    "L1715_Sim16_Notecards for HCP Evaluation 1.docx": "Sim16_Notecards for HCP Evaluation 1.docx",
    "L1715_Sim16_Notecards for HCP Evaluation 2.docx": "Sim16_Notecards for HCP Evaluation 2.docx",
    "L1715_Sim22 SCAT6 Rain Shoemaker 96 Hours.pdf": "Sim22 SCAT6 Rain Shoemaker 96 Hours.pdf",
}
LEGACY_DOWNLOAD_NAMES = {
    22: {
        "scat6-rain-shoemaker-96-hours.pdf": "Sim22 SCAT6 Rain Shoemaker 96 Hours.pdf",
    },
}
ADDITIONAL_DOWNLOADS = {
    15: (("Rubric: Mountain Biking", "Rubric-mountain-biking.docx"),),
}

# Composite proctor packets requested by the production notes in each manuscript.
# Component filenames refer to the individual downloads generated immediately
# before the packet is assembled.
PROCTOR_PACKETS = {
    1: (
        ("Information for Proctor for Grade I Sprain", "information-for-proctor-for-grade-i-sprain.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-1.docx", "instructions-for-embedded-personnel.docx")),
        ("Information for Proctor for Grade II Sprain", "information-for-proctor-for-grade-ii-sprain.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-2.docx", "instructions-for-embedded-personnel.docx")),
    ),
    2: (("Information for Proctor for Exercise Illness", "information-for-proctor-for-exercise-illness.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-standardized-patient.docx")),),
    3: (("Information for Proctor", "information-for-proctor.docx", ("rubric.docx", "instructions-for-football-sp.docx", "instructions-for-swimming-sp.docx", "instructions-for-fan-in-the-stands-sp.docx", "instructions-for-coach-heart-attack-in-the-stairwell-sp.docx", "instructions-for-football-proctor.docx", "instructions-for-swimming-proctor.docx", "instructions-for-fan-in-the-stands-proctor.docx", "instructions-for-coach-in-the-stairwell-proctor.docx", "instructions-for-ems-all-situations.docx")),),
    4: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "rubric.docx")),),
    5: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "rubric.docx")),),
    6: (
        ("Information for Proctor for Asthma SP", "information-for-proctor-for-asthma-sp.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-asthma.docx")),
        ("Information for Proctor for Arrhythmia SP", "information-for-proctor-for-arrhythmia-sp.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-arrhythmia.docx")),
        ("Information for Proctor for Anxiety SP", "information-for-proctor-for-anxiety-sp.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-anxiety.docx")),
        ("Information for Proctor for Anaphylaxis SP", "information-for-proctor-for-anaphylaxis-sp.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-anaphylaxis.docx")),
    ),
    7: (("Information for Proctor", "information-for-proctor.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-ep-coach.docx")),),
    8: (
        ("Information for Proctor for SP 1 Spleen", "information-for-proctor-for-sp-1-spleen.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-1-spleen.docx")),
        ("Information for Proctor for SP 2 Liver", "information-for-proctor-for-sp-2-liver.docx", ("rubric.docx", "situation-and-setting.docx", "instructions-for-sp-2-liver.docx")),
    ),
    9: (("Information for Proctor", "information-for-proctor.docx", ("rubric.docx", "situation-and-setting.docx", "embedded-personnel-coach.docx", "athlete-summary-sickle-cell.docx", "athlete-summary-asthma.docx", "athlete-summary-epilepsy.docx")),),
    10: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "rubric.docx", "instructions-for-sp-1-hotel-room-gi.docx", "instructions-for-sp-2-sideline-pool-deck.docx", "instructions-for-sp-3-athletic-training-clinic-dance-department-college.docx", "instructions-for-sp-4-athletic-training-clinic-high-school.docx", "instructions-for-sp-5-track-sideline.docx")),),
    12: tuple(
        (f"Information for Proctor for Station {station}", f"information-for-proctor-for-station-{station}.docx", ("rubric.docx", "situation-and-setting.docx", f"sp-information-for-station-{station}.docx"))
        for station in range(1, 7)
    ),
    13: (("Information for Proctor", "information-for-proctor.docx", ("instructions-for-sp.docx", "rubric.docx", "situation-and-setting.docx")),),
    15: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-hockey-sp.docx", "instructions-for-unconscious-basketball-sp-and-embedded-personnel-ep.docx", "instructions-for-mountain-biking-sp-and-ep.docx", "instructions-for-football-sp.docx", "rubric-hockey.docx", "rubric-unconscious-basketball.docx", "Rubric-mountain-biking.docx", "rubric-football.docx")),),
    18: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp.docx", "instructions-for-ep-proctor.docx", "instructions-for-ep-ems.docx", "instructions-for-ep-parent.docx", "event-medical-form.docx", "rubric.docx")),),
    19: (("Information for Proctor", "information-for-proctor.docx", ("instructions-for-sp.docx", "situation-and-setting.docx", "rubric.docx")),),
    20: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp-patient-a.docx", "instructions-for-sp-patient-b.docx", "rubric.docx")),),
    21: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp-patient-a.docx", "instructions-for-sp-patient-b.docx", "rubric.docx")),),
    22: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-ep.docx", "rubric.docx")),),
    23: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp-1-and-2.docx", "rubric-ehs-station-1.docx", "rubric-ehs-station-2.docx")),),
    24: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp.docx", "rubric.docx")),),
    25: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp.docx", "rubric.docx")),),
    26: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp.docx", "rubric.docx")),),
    27: (("Information for Proctor", "information-for-proctor.docx", ("instructions-for-all-sps.docx", "instructions-for-ep.docx", "situation-and-setting.docx", "rubric.docx")),),
    28: (("Information for Proctor", "information-for-proctor.docx", ("instructions-for-all-sps.docx", "instructions-for-all-eps.docx", "situation-and-setting.docx", "rubric.docx")),),
    29: (("Information for Proctor", "information-for-proctor.docx", ("instructions-for-all-sps.docx", "instructions-for-ep.docx", "rubric.docx")),),
    30: (("Information for Proctor", "information-for-proctor.docx", ("instructions-for-sp.docx", "situation-and-setting.docx", "rubric.docx")),),
    31: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-all-eps.docx", "rubric.docx")),),
    32: (("Information for Proctor", "information-for-proctor.docx", ("situation-and-setting.docx", "instructions-for-sp.docx", "rubric.docx")),),
}


def manuscript_sources():
    """Return the single primary instructor manuscript for each simulation."""
    sources = {}
    supporting_names = {
        filename
        for filenames in SUPPORTING_RESOURCES.values()
        for filename in filenames
    }
    for path in sorted(MANUSCRIPTS.glob("L1715_Sim*.docx")):
        match = SIM_RE.match(path.name)
        if not match or path.name in supporting_names:
            continue
        number = int(match.group(1))
        if number in sources:
            raise ValueError(
                f"Multiple primary manuscripts found for simulation {number}: "
                f"{sources[number].name}, {path.name}"
            )
        sources[number] = path
    return [sources[number] for number in sorted(sources)]


def text_of(node):
    return "".join(node.xpath(".//w:t/text()", namespaces=NS))


def remove_prefix(node, prefix):
    remaining = prefix
    for text_node in node.xpath(".//w:t", namespaces=NS):
        value = text_node.text or ""
        if not remaining:
            break
        count = min(len(value), len(remaining))
        if value[:count] != remaining[:count]:
            return
        text_node.text = value[count:]
        remaining = remaining[count:]


def remove_suffix(node, suffix):
    remaining = suffix
    for text_node in reversed(node.xpath(".//w:t", namespaces=NS)):
        value = text_node.text or ""
        if not remaining:
            break
        count = min(len(value), len(remaining))
        if value[-count:] != remaining[-count:]:
            return
        text_node.text = value[:-count] if count else value
        remaining = remaining[:-count]


def prepare_reader_node(node):
    """Clone a manuscript block and remove production-only prefixes and edge space."""
    clone = deepcopy(node)
    value = text_of(clone)
    leading = re.match(r"^\s*", value).group(0)
    if leading:
        remove_prefix(clone, leading)
    value = text_of(clone)
    photo_prefix = re.match(r"^(?:PHOTO HERE\s*)+", value)
    if photo_prefix:
        remove_prefix(clone, photo_prefix.group(0))
    value = text_of(clone)
    marker = PRODUCTION_PREFIX.match(value)
    if marker:
        remove_prefix(clone, marker.group(0))
    value = text_of(clone)
    trailing = re.search(r"\s*$", value).group(0)
    if trailing:
        remove_suffix(clone, trailing)
    return clone


def body_children(path):
    with ZipFile(path) as archive:
        root = etree.fromstring(
            accept_tracked_changes_xml(archive.read("word/document.xml"))
        )
    return root, root.find("w:body", NS)


def body_blocks(parent):
    """Yield block-level content, including blocks inside Word content controls."""
    for child in parent:
        if child.tag in {qn("w:p"), qn("w:tbl"), qn("w:sectPr")}:
            yield child
        elif child.tag in {qn("w:sdt"), qn("w:sdtContent"), qn("w:customXml")}:
            yield from body_blocks(child)


def document_formatting(path):
    """Load hyperlink targets and effective character-style properties."""
    with ZipFile(path) as archive:
        relationships = {}
        rels_name = "word/_rels/document.xml.rels"
        if rels_name in archive.namelist():
            rels_root = etree.fromstring(archive.read(rels_name))
            relationships = {
                relation.get("Id"): relation.get("Target")
                for relation in rels_root
                if relation.get("Id") and relation.get("Target")
            }

        raw_styles = {}
        if "word/styles.xml" in archive.namelist():
            styles_root = etree.fromstring(archive.read("word/styles.xml"))
            for style in styles_root.findall("w:style", NS):
                identifier = style.get(qn("w:styleId"))
                if not identifier:
                    continue
                based_on = style.find("w:basedOn", NS)
                raw_styles[identifier] = {
                    "based_on": based_on.get(qn("w:val")) if based_on is not None else None,
                    "properties": formatting_properties(style.find("w:rPr", NS)),
                }

    resolved = {}

    def resolve(identifier, seen=None):
        if identifier in resolved:
            return resolved[identifier]
        if identifier not in raw_styles:
            return {}
        seen = set() if seen is None else seen
        if identifier in seen:
            return {}
        seen.add(identifier)
        entry = raw_styles[identifier]
        properties = dict(resolve(entry["based_on"], seen)) if entry["based_on"] else {}
        properties.update(entry["properties"])
        resolved[identifier] = properties
        return properties

    for identifier in raw_styles:
        resolve(identifier)
    return relationships, resolved


def formatting_properties(run_properties):
    if run_properties is None:
        return {}
    result = {}
    for name in ("b", "i", "strike", "caps", "smallCaps"):
        element = run_properties.find(f"w:{name}", NS)
        if element is not None:
            result[name] = element.get(qn("w:val"), "1").lower() not in {
                "0", "false", "off", "none"
            }
    underline = run_properties.find("w:u", NS)
    if underline is not None:
        result["u"] = underline.get(qn("w:val"), "single").lower() not in {
            "0", "false", "off", "none"
        }
    vertical = run_properties.find("w:vertAlign", NS)
    if vertical is not None:
        result["vertAlign"] = vertical.get(qn("w:val"), "baseline")
    return result


def effective_run_properties(run, styles):
    run_properties = run.find("w:rPr", NS)
    result = {}
    if run_properties is not None:
        style = run_properties.find("w:rStyle", NS)
        if style is not None:
            result.update(styles.get(style.get(qn("w:val")), {}))
        result.update(formatting_properties(run_properties))
    return result


def run_html(run, styles):
    parts = []
    for child in run:
        if child.tag == qn("w:t"):
            parts.append(escape(child.text or ""))
        elif child.tag in {qn("w:br"), qn("w:cr")}:
            parts.append("<br>")
        elif child.tag == qn("w:tab"):
            parts.append('<span class="manuscript-tab" aria-hidden="true"></span>')
    rendered = "".join(parts)
    if not rendered:
        return ""
    properties = effective_run_properties(run, styles)
    wrappers = []
    if properties.get("b"):
        wrappers.append("strong")
    if properties.get("i"):
        wrappers.append("em")
    if properties.get("u"):
        wrappers.append("u")
    if properties.get("strike"):
        wrappers.append("s")
    if properties.get("vertAlign") in {"superscript", "subscript"}:
        wrappers.append("sup" if properties["vertAlign"] == "superscript" else "sub")
    for tag in reversed(wrappers):
        rendered = f"<{tag}>{rendered}</{tag}>"
    if properties.get("caps"):
        rendered = f'<span class="manuscript-caps">{rendered}</span>'
    if properties.get("smallCaps"):
        rendered = f'<span class="manuscript-small-caps">{rendered}</span>'
    return rendered


def inline_html(node, relationships, styles):
    def render_children(parent):
        output = []
        for child in parent:
            if child.tag == qn("w:r"):
                output.append(run_html(child, styles))
            elif child.tag == qn("w:hyperlink"):
                content = render_children(child)
                relation_id = child.get(R + "id")
                target = relationships.get(relation_id)
                anchor = child.get(qn("w:anchor"))
                href = target or (f"#{anchor}" if anchor else "")
                output.append(
                    f'<a href="{escape(href, quote=True)}">{content}</a>' if href else content
                )
            elif child.tag not in {qn("w:del"), qn("w:instrText")}:
                output.append(render_children(child))
        return "".join(output)

    return render_children(node)


def next_nonempty_text(children, start):
    for node in children[start:]:
        value = text_of(node).strip()
        if value:
            return value
    return ""


def download_sections(path):
    _root, body = body_children(path)
    children = list(body_blocks(body))
    sections = []
    for index, node in enumerate(children):
        marker = text_of(node).strip()
        if not BEGIN_RE.search(marker):
            continue
        end = next(
            (candidate for candidate in range(index + 1, len(children))
             if text_of(children[candidate]).strip().startswith("\\qqEND downloadable content")
             or BEGIN_RE.search(text_of(children[candidate]))),
            len(children),
        )
        selected = [deepcopy(item) for item in children[index + 1:end]]
        if "<title>" in marker:
            inline = deepcopy(node)
            remove_prefix(inline, marker.split("<title>", 1)[0])
            selected.insert(0, inline)
        title_match = next(
            (
                match
                for item in selected
                if (match := re.match(r"<(?:title|b|a)>(.*)", text_of(item).strip()))
            ),
            None,
        )
        button_match = BUTTON_RE.search(marker.rstrip("\\"))
        title = title_match.group(1).strip() if title_match else ""
        if button_match:
            button = button_match.group(1).strip().rstrip("\\")
        else:
            button = marker.split("BEGIN downloadable content", 1)[1].strip(" \\.:")
            button = button or title
        if not title:
            title = button
        if title and selected:
            sections.append({"button": button or title, "title": title, "nodes": selected})
    return sections


def clean_download_nodes(nodes):
    cleaned = []
    in_note = False
    pending_page_break = False
    for node in nodes:
        value = text_of(node).strip()
        if PAGE_BREAK_MARKER_RE.match(value):
            pending_page_break = True
            continue
        if value.startswith("\\qqID:") or value.startswith("\\qqPSM:"):
            in_note = not NOTE_END_RE.search(value)
            continue
        if in_note:
            if NOTE_END_RE.search(value):
                in_note = False
                continue
            # A new reader-facing block means the preceding production note
            # was a single unclosed instruction, not a multiline note.
            if not PRODUCTION_PREFIX.match(value):
                continue
            in_note = False
        if value.startswith("\\qqINSERT"):
            if image_for_marker(value) is None:
                continue
        if value.startswith("\\qq") and not value.startswith("\\qqINSERT"):
            continue
        clone = deepcopy(node)
        clone_text = text_of(clone)
        match = PRODUCTION_PREFIX.match(clone_text)
        if match:
            remove_prefix(clone, match.group(0))
        if pending_page_break and clone.tag == qn("w:p"):
            paragraph_properties = clone.find("w:pPr", NS)
            if paragraph_properties is None:
                paragraph_properties = etree.Element(qn("w:pPr"))
                clone.insert(0, paragraph_properties)
            if paragraph_properties.find("w:pageBreakBefore", NS) is None:
                paragraph_properties.append(etree.Element(qn("w:pageBreakBefore")))
            pending_page_break = False
        if text_of(clone).strip() or clone.tag == qn("w:tbl"):
            cleaned.append(clone)
    return cleaned


def document_for_nodes(source, nodes):
    with ZipFile(source) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    body = root.find("w:body", NS)
    section_properties = body.find("w:sectPr", NS)
    for node in list(body):
        body.remove(node)
    for node in clean_download_nodes(nodes):
        body.append(node)
    if section_properties is not None:
        body.append(deepcopy(section_properties))
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def slug(value):
    value = value.lower().replace("’", "").replace("'", "")
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value[:90] or "instructor-download"


def download_specs(source):
    results = []
    used = set()
    for section in download_sections(source):
        if section["button"] == "Information for Proctor":
            # These markers contain production assembly directions, not reader-facing copy.
            continue
        nodes = clean_download_nodes(section["nodes"])
        meaningful = [text_of(node).strip() for node in nodes if text_of(node).strip()]
        if len(meaningful) <= 1:
            continue
        base = slug(section["button"])
        filename = f"{base}.docx"
        suffix = 2
        while filename in used:
            filename = f"{base}-{suffix}.docx"
            suffix += 1
        used.add(filename)
        results.append((section, section["button"], filename))
    return results


def download_manifest(source, sim_number=None):
    results = [(label, filename) for _section, label, filename in download_specs(source)]
    results.extend(
        (label, filename)
        for label, filename, _components in PROCTOR_PACKETS.get(sim_number, ())
    )
    results.extend(ADDITIONAL_DOWNLOADS.get(sim_number, ()))
    return results


def write_composite_docx(source, output, title, component_paths):
    """Combine generated downloads into one page-separated proctor packet."""
    document = Document(source)
    body = document._element.body
    section_properties = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not section_properties:
            body.remove(child)

    title_paragraph = document.add_paragraph(title)

    for component_index, component_path in enumerate(component_paths):
        if not component_path.exists():
            raise FileNotFoundError(
                f"Missing proctor-packet component: {component_path.relative_to(ROOT)}"
            )
        if component_index:
            separator = document.add_paragraph()
            separator.add_run().add_break(WD_BREAK.PAGE)
        component = Document(component_path)
        for child in component._element.body:
            if child.tag == qn("w:sectPr"):
                continue
            insertion_index = len(body) - 1 if section_properties is not None else len(body)
            body.insert(insertion_index, deepcopy(child))

    document.core_properties.title = title
    document.save(output)
    finalize_document(output)


def write_proctor_packets(source, sim_number, output_dir):
    results = []
    for label, filename, components in PROCTOR_PACKETS.get(sim_number, ()):
        output = output_dir / filename
        write_composite_docx(
            source,
            output,
            label,
            [output_dir / component for component in components],
        )
        results.append((label, filename))
    return results


def write_downloads(source, sim_number):
    output_dir = DOWNLOADS / f"simulation-{sim_number}"
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for section, label, filename in download_specs(source):
        output = output_dir / filename
        write_docx(source, output, document_for_nodes(source, section["nodes"]))
        finalize_document(output)
        results.append((label, filename))
    supporting = [
        MANUSCRIPTS / filename
        for filename in SUPPORTING_RESOURCES.get(sim_number, ())
        if (MANUSCRIPTS / filename).exists()
    ]
    packaged_supporting = []
    for path in supporting:
        output_name = SUPPORTING_OUTPUT_NAMES.get(path.name, path.name)
        if path.suffix.lower() == ".docx":
            packaged = output_dir / output_name
            copy2(path, packaged)
            stamp_footer(packaged)
            packaged_supporting.append(packaged)
        elif output_name != path.name:
            packaged = output_dir / output_name
            if not packaged.exists():
                copy2(path, packaged)
            packaged_supporting.append(packaged)
        else:
            packaged_supporting.append(path)
    additional = [
        (label, filename)
        for label, filename in ADDITIONAL_DOWNLOADS.get(sim_number, ())
        if (output_dir / filename).exists()
    ]
    proctor_packets = write_proctor_packets(source, sim_number, output_dir)
    if results or supporting or additional or proctor_packets:
        zip_path = output_dir / "all-instructor-downloads.zip"
        with ZipFile(zip_path, "w", ZIP_DEFLATED) as archive:
            for _label, filename in results:
                archive.write(output_dir / filename, filename)
            for _label, filename in proctor_packets:
                archive.write(output_dir / filename, filename)
            for _label, filename in additional:
                archive.write(output_dir / filename, filename)
            for path in packaged_supporting:
                archive.write(path, path.name)
    return results + proctor_packets + additional


def cell_html(cell, relationships, styles):
    parts = []
    for paragraph in cell.findall("w:p", NS):
        clean_paragraph = prepare_reader_node(paragraph)
        value = text_of(clean_paragraph)
        if value:
            parts.append(inline_html(clean_paragraph, relationships, styles))
    return "<br>".join(parts)


def table_html(node, relationships, styles):
    rows = []
    for row_index, row in enumerate(node.findall("w:tr", NS)):
        cells = row.findall("w:tc", NS)
        if row_index == 0:
            markup = "".join(
                f'<th scope="col" data-manuscript-block>{cell_html(cell, relationships, styles)}</th>'
                for cell in cells
            )
        else:
            markup = "".join(
                f"<td data-manuscript-block>{cell_html(cell, relationships, styles)}</td>"
                for cell in cells
            )
        rows.append("<tr>" + markup + "</tr>")
    header = rows[0] if rows else ""
    body_rows = "".join(rows[1:])
    return (
        '<div class="table-scroll" role="region" aria-label="Scrollable data table" tabindex="0">'
        '<table class="manuscript-table"><thead>' + header + "</thead><tbody>"
        + body_rows + "</tbody></table></div>"
    )


def image_for_marker(value):
    marker = re.match(r"^\\qqINSERT:?\s+(.*)$", value)
    if not marker or not marker.group(1).strip():
        return None
    stem = PureWindowsPath(marker.group(1).strip()).stem
    for suffix in (".png", ".jpg", ".jpeg"):
        candidate = ROOT / "assets" / "images" / f"{stem}{suffix}"
        if candidate.exists():
            return candidate
    return None


def main_content(path):
    _root, body = body_children(path)
    relationships, styles = document_formatting(path)
    output = []
    in_download = False
    in_note = False
    active_list = None
    inline_numbered_active = False
    sim_match = SIM_RE.match(path.name)
    sim_number = int(sim_match.group(1)) if sim_match else None
    inline_numbered_items = INLINE_NUMBERED_ITEMS.get(sim_number, frozenset())

    def close_list():
        nonlocal active_list, inline_numbered_active
        if active_list:
            output.append(f"</{active_list}>")
            active_list = None
        inline_numbered_active = False

    for node in body_blocks(body):
        value = text_of(node).strip()
        value = re.sub(r"^(?:PHOTO HERE\s*)+", "", value).strip()
        if BEGIN_RE.search(value):
            close_list()
            in_download = True
            in_note = False
            continue
        if in_download:
            if value.startswith("\\qqEND downloadable content"):
                in_download = False
                continue
            # Supporting-file button markers sometimes omit an END marker.
            # A top-level manuscript heading is an unambiguous return to the
            # web-page content; downloadable sections use <title> headings.
            if not value.startswith("<a>"):
                continue
            in_download = False
        if value.startswith("\\qqID:") or value.startswith("\\qqPSM:"):
            in_note = not NOTE_END_RE.search(value)
            continue
        if in_note:
            if NOTE_END_RE.search(value):
                in_note = False
                continue
            # Do not let a malformed or unclosed production note consume a
            # subsequent manuscript section.
            if not PRODUCTION_PREFIX.match(value):
                continue
            in_note = False
        if value.startswith("\\qqINSERT"):
            close_list()
            image = image_for_marker(value)
            if image:
                output.append(
                    '<figure class="manuscript-figure"><img src="../assets/images/'
                    f'{escape(image.name)}" alt=""></figure>'
                )
            continue
        if not value or value.startswith("Navigation menu/button") or value.startswith("\\qq"):
            continue
        if node.tag == qn("w:tbl"):
            close_list()
            output.append(table_html(node, relationships, styles))
            continue
        clean_node = prepare_reader_node(node)
        clean = text_of(clean_node)
        marker = PRODUCTION_PREFIX.match(value)
        tag_name = marker.group(1) if marker else ""
        if tag_name in ("cn", "ct"):
            continue
        html_tag = HEADING_TAGS.get(tag_name, "p")
        if clean in FORCED_TOP_LEVEL_HEADINGS.get(sim_number, frozenset()):
            html_tag = "h2"
        if html_tag.startswith("h"):
            close_list()
            rendered = inline_html(clean_node, relationships, styles)
            output.append(f"<{html_tag} data-manuscript-block>{rendered}</{html_tag}>")
            list_config = LIST_SECTIONS.get(clean)
            if list_config:
                active_list, list_class = list_config
                output.append(f'<{active_list} class="manuscript-list {list_class}">')
            continue
        if html_tag == "p" and clean in inline_numbered_items:
            if not inline_numbered_active:
                close_list()
                active_list = "ol"
                inline_numbered_active = True
                output.append('<ol class="manuscript-list numbered-notes-list">')
            rendered = inline_html(clean_node, relationships, styles)
            output.append(f"<li data-manuscript-block>{rendered}</li>")
            continue
        if inline_numbered_active:
            close_list()
        if active_list and html_tag == "p":
            number_match = SOURCE_LIST_NUMBER_RE.match(clean)
            if number_match:
                prefix = number_match.group(1)
                remove_prefix(clean_node, prefix)
                rendered = (
                    f'<span class="source-list-number">{escape(prefix)}</span>'
                    f'{inline_html(clean_node, relationships, styles)}'
                )
            else:
                rendered = inline_html(clean_node, relationships, styles)
            output.append(f"<li data-manuscript-block>{rendered}</li>")
            continue
        close_list()
        css = ' class="figure-caption"' if tag_name == "fc" else ""
        rendered = inline_html(clean_node, relationships, styles)
        for source_text, replacement_text in PAGE_TEXT_REPLACEMENTS.get(
            sim_number, {}
        ).items():
            rendered = rendered.replace(source_text, replacement_text)
        for url in SIMULATION_LINKS.get(sim_number, ()):
            escaped_url = escape(url)
            if url in rendered and f'href="{escaped_url}"' not in rendered:
                rendered = rendered.replace(
                    url,
                    f'<a href="{escaped_url}" target="_blank" rel="noopener noreferrer">{escaped_url}</a>',
                )
        output.append(f"<{html_tag}{css} data-manuscript-block>{rendered}</{html_tag}>")
    close_list()
    return output


def doc_title(path, fallback):
    _root, body = body_children(path)
    for paragraph in body_blocks(body):
        if paragraph.tag != qn("w:p"):
            continue
        value = text_of(paragraph).strip()
        if value.startswith("<ct>"):
            title = re.sub(r"^Simulation \d+:?\s*", "", value.removeprefix("<ct>"))
            return title or fallback
    return fallback


def page_shell(title, kicker, content):
    return f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)}</title><link rel="stylesheet" href="page.css"></head>
<body data-manuscript><header class="lesson-header"><p class="kicker">{escape(kicker)}</p><h1>{escape(title)}</h1></header>
<main class="page">{content}</main></body></html>
'''


def link_debriefing_methods(rendered):
    """Apply the manuscript's production link instruction within the Debrief card."""
    section_pattern = re.compile(
        r'(<section\b[^>]*>\s*<h2\b[^>]*>Debrief</h2>)(.*?)(</section>)',
        re.IGNORECASE | re.DOTALL,
    )

    def link_reference(match):
        body = match.group(2)
        if 'href="debriefing-methods.html"' in body:
            return match.group(0)
        reference_pattern = re.compile(r"\bdebriefing methods?\b", re.IGNORECASE)
        linked_body, count = reference_pattern.subn(
            lambda reference: (
                '<a class="manuscript-instruction-link" '
                f'href="debriefing-methods.html">{reference.group(0)}</a>'
            ),
            body,
            count=1,
        )
        if count == 0:
            return match.group(0)
        return match.group(1) + linked_body + match.group(3)

    updated, count = section_pattern.subn(link_reference, rendered, count=1)
    if count != 1:
        raise ValueError("Simulation page is missing its Debrief section")
    return updated


def promote_existing_page_sections(page, sim_number):
    """Promote configured nested headings into their own generated cards."""
    rendered = page.read_text(encoding="utf-8")
    updated = rendered
    for heading in FORCED_TOP_LEVEL_HEADINGS.get(sim_number, frozenset()):
        for source_tag in ("h3", "h4", "h5"):
            needle = f'<{source_tag} data-manuscript-block>{escape(heading)}</{source_tag}>'
            if needle not in updated:
                continue
            replacement = (
                '</section><section class="content-card">'
                f'<h2 data-manuscript-block>{escape(heading)}</h2>'
            )
            updated = updated.replace(needle, replacement, 1)
            break
    if updated != rendered:
        page.write_text(updated, encoding="utf-8")


def update_existing_supporting_links(page, sim_number):
    """Align generated page links with renamed packaged supporting files."""
    rendered = page.read_text(encoding="utf-8")
    updated = rendered
    for source_name in SUPPORTING_RESOURCES.get(sim_number, ()):
        output_name = SUPPORTING_OUTPUT_NAMES.get(source_name)
        if not output_name:
            continue
        target = f"../assets/downloads/simulation-{sim_number}/{output_name}"
        updated = updated.replace(f"../assets/Manuscripts/{source_name}", target)
        updated = updated.replace(
            f"../assets/downloads/simulation-{sim_number}/{source_name}", target
        )
    for legacy_name, output_name in LEGACY_DOWNLOAD_NAMES.get(sim_number, {}).items():
        updated = updated.replace(
            f"../assets/downloads/simulation-{sim_number}/{legacy_name}",
            f"../assets/downloads/simulation-{sim_number}/{output_name}",
        )
    if updated != rendered:
        page.write_text(updated, encoding="utf-8")


def build_manuscript_page(source, number, fallback_title, write_download_assets=True):
    title = doc_title(source, fallback_title)
    blocks = main_content(source)
    for heading, paragraphs in SECTION_CONTENT_ADDITIONS.get(number, {}).items():
        heading_block = f"<h2 data-manuscript-block>{escape(heading)}</h2>"
        try:
            heading_index = blocks.index(heading_block)
        except ValueError as error:
            raise ValueError(
                f"Simulation {number} is missing the {heading!r} section"
            ) from error
        blocks[heading_index + 1 : heading_index + 1] = [
            f"<p data-manuscript-block>{escape(paragraph)}</p>"
            for paragraph in paragraphs
        ]
    cards = []
    current = []
    for block in blocks:
        if block.startswith("<h2") and current:
            cards.append('<section class="content-card">' + "".join(current) + "</section>")
            current = []
        current.append(block)
    if current:
        cards.append('<section class="content-card">' + "".join(current) + "</section>")
    downloads = (
        write_downloads(source, number)
        if write_download_assets
        else download_manifest(source, number)
    )
    supporting = [
        filename
        for filename in SUPPORTING_RESOURCES.get(number, ())
        if (MANUSCRIPTS / filename).exists()
    ]
    if downloads or supporting:
        links = ['<a class="download-card featured" href="../assets/downloads/simulation-{0}/all-instructor-downloads.zip" download><span class="file-icon" aria-hidden="true">ZIP</span><span><strong>All Instructor Downloads</strong><small>ZIP archive</small></span><span class="download-arrow" aria-hidden="true">↓</span></a>'.format(number)]
        for label, filename in downloads:
            links.append(f'<a class="download-card" href="../assets/downloads/simulation-{number}/{escape(filename)}" download><span class="file-icon" aria-hidden="true">DOCX</span><span><strong>{escape(label)}</strong><small>Word document</small></span><span class="download-arrow" aria-hidden="true">↓</span></a>')
        for filename in supporting:
            extension = Path(filename).suffix.removeprefix(".").upper()
            label = re.sub(r"^L1715_Sim\d{2}[_ ]*", "", Path(filename).stem).replace("_", " ")
            output_filename = SUPPORTING_OUTPUT_NAMES.get(filename, filename)
            href = (
                f"../assets/downloads/simulation-{number}/{output_filename}"
                if extension == "DOCX" or output_filename != filename
                else f"../assets/Manuscripts/{filename}"
            )
            links.append(
                f'<a class="download-card" href="{escape(href)}" download>'
                f'<span class="file-icon" aria-hidden="true">{escape(extension)}</span><span><strong>{escape(label)}</strong>'
                f'<small>Supporting instructor resource</small></span><span class="download-arrow" aria-hidden="true">&darr;</span></a>'
            )
        cards.append('<section class="content-card"><h2>Instructor Downloads</h2><div class="download-grid activity-grid">' + "".join(links) + "</div></section>")
    page = PAGES / f"simulation-{number}.html"
    rendered = page_shell(title, f"Simulation {number}", "".join(cards))
    page.write_text(link_debriefing_methods(rendered), encoding="utf-8")
    promote_existing_page_sections(page, number)
    update_existing_supporting_links(page, number)
    return title, len(downloads)


def build_introduction(write_download_assets=True):
    source = MANUSCRIPTS / "L1715_Introduction.docx"
    if write_download_assets:
        stamp_footer(DOWNLOADS / "copyright-page-placeholder.docx")
    page_path = PAGES / "introduction.html"
    try:
        blocks = main_content(source)
    except PermissionError:
        # OneDrive or Word can temporarily lock a hydrated manuscript. Keep the
        # last verified page usable and attach the navigation enhancement.
        rendered = page_path.read_text(encoding="utf-8")
        if "introduction-nav.js" not in rendered:
            rendered = rendered.replace(
                "</body>", '<script src="introduction-nav.js"></script></body>'
            )
            page_path.write_text(rendered, encoding="utf-8")
        print(f"LOCKED: retained {page_path.relative_to(ROOT)} and added section navigation")
        return
    headings = []
    anchored_blocks = []
    for block in blocks:
        match = re.match(r'(<h2)([^>]*>)(.*?)(</h2>)$', block)
        if not match:
            anchored_blocks.append(block)
            continue
        label = unescape(re.sub(r"<[^>]+>", "", match.group(3)))
        anchor = f"introduction-{slug(label)}"
        headings.append((label, anchor))
        anchored_blocks.append(
            f'{match.group(1)} id="{anchor}"{match.group(2)}{match.group(3)}{match.group(4)}'
        )

    jump_links = "".join(
        f'<a class="section-jump" href="#{anchor}">{escape(label)}</a>'
        for label, anchor in headings
    )
    jump_links += '<a class="section-jump" href="#introduction-copyright">Copyright</a>'
    navigation = (
        '<details class="content-card introduction-navigation" id="introduction-navigation" open>'
        '<summary>Section navigation</summary>'
        f'<div class="section-jump-grid">{jump_links}</div></details>'
    )

    cards, current = [], []
    for block in anchored_blocks:
        if block.startswith("<h2") and current:
            cards.append('<section class="content-card">' + "".join(current) + "</section>")
            current = []
        current.append(block)
    if current:
        cards.append('<section class="content-card">' + "".join(current) + "</section>")
    for index in range(1, len(cards)):
        cards[index] = cards[index].replace(
            "</section>",
            '<a class="section-return" href="#introduction-navigation">↑ Return to section navigation</a></section>',
            1,
        )
    copyright_card = (
        '<section class="content-card" id="introduction-copyright"><h2>Copyright</h2>'
        '<div class="download-grid activity-grid"><a class="download-card featured" '
        'href="../assets/downloads/copyright-page-placeholder.docx" download>'
        '<span class="file-icon" aria-hidden="true">DOCX</span><span>'
        '<strong>Download Copyright Page</strong><small>Placeholder Word document</small>'
        '</span><span class="download-arrow" aria-hidden="true">↓</span></a></div>'
        '<a class="section-return" href="#introduction-navigation">'
        '↑ Return to section navigation</a></section>'
    )
    rendered = page_shell(
        "Introduction", "Instructor guide", navigation + "".join(cards) + copyright_card
    )
    rendered = rendered.replace("</body>", '<script src="introduction-nav.js"></script></body>')
    page_path.write_text(rendered, encoding="utf-8")


def build_debriefing():
    source = MANUSCRIPTS / "L1715_Debriefing Methods.docx"
    blocks = [
        block
        for block in main_content(source)
        if unescape(re.sub(r"<[^>]+>", "", block)) != "Debriefing Methods"
    ]
    content = '<section class="content-card">' + "".join(blocks) + "</section>"
    (PAGES / "debriefing-methods.html").write_text(page_shell("Debriefing Methods", "Instructor resource", content), encoding="utf-8")


def resource_page(filename, title, label):
    content = f'<section class="content-card"><div class="download-grid"><a class="download-card featured" href="../assets/Manuscripts/{escape(filename)}" download><span class="file-icon" aria-hidden="true">{escape(label)}</span><span><strong>{escape(title)}</strong><small>{escape(filename)}</small></span><span class="download-arrow" aria-hidden="true">↓</span></a></div></section>'
    return page_shell(title, "Instructor resource", content)


def build_resource_pages():
    resources = (
        ("simulation-checklist.html", "L1715_Simulation Checklist.xlsx", "Simulation Checklist", "XLSX"),
        ("simulation-finder.html", "L1715_Simulation Finder.xlsx", "Simulation Finder", "XLSX"),
    )
    for page, source, title, label in resources:
        (PAGES / page).write_text(resource_page(source, title, label), encoding="utf-8")


def navigation_titles():
    source = (ROOT / "data" / "navigation.js").read_text(encoding="utf-8")
    return {
        int(number): title
        for number, title in re.findall(
            r'id: "simulation-(\d+)", title: "Simulation \d+: ([^"]+)"', source
        )
    }


def build_pending_pages(available):
    titles = navigation_titles()
    for number in range(1, 33):
        if number in available:
            continue
        title = titles.get(number, f"Simulation {number}")
        resources = SUPPORTING_RESOURCES.get(number, ())
        links = []
        for filename in resources:
            extension = Path(filename).suffix.removeprefix(".").upper()
            links.append(
                f'<a class="download-card" href="../assets/Manuscripts/{escape(filename)}" download>'
                f'<span class="file-icon" aria-hidden="true">{escape(extension)}</span><span><strong>{escape(filename)}</strong>'
                f'<small>Supporting instructor resource</small></span><span class="download-arrow" aria-hidden="true">↓</span></a>'
            )
        resource_section = ""
        if links:
            resource_section = '<section class="content-card"><h2>Available Instructor Resources</h2><div class="download-grid">' + "".join(links) + "</div></section>"
        notice = '<section class="content-card tint"><h2>Instructor Manuscript Status</h2><p>The instructor manuscript for this simulation is not yet available.</p></section>'
        (PAGES / f"simulation-{number}.html").write_text(
            page_shell(title, f"Simulation {number}", notice + resource_section), encoding="utf-8"
        )


def main():
    pages_only = "--pages-only" in sys.argv[1:]
    requested = {int(arg) for arg in sys.argv[1:] if arg != "--pages-only"}
    titles = navigation_titles()
    build_introduction(write_download_assets=not pages_only)
    build_debriefing()
    build_resource_pages()
    built = []
    for source in manuscript_sources():
        match = SIM_RE.match(source.name)
        if not match:
            continue
        number = int(match.group(1))
        if requested and number not in requested:
            continue
        fallback = TITLE_OVERRIDES.get(number, titles.get(number, f"Simulation {number}"))
        title, count = build_manuscript_page(
            source, number, fallback, write_download_assets=not pages_only
        )
        built.append(number)
        print(f"BUILT: simulation-{number} ({title}; {count} downloads)")
    if not requested:
        build_pending_pages(set(built))
    print("AVAILABLE:", ", ".join(map(str, built)))


if __name__ == "__main__":
    main()
